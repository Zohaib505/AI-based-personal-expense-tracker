
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def add_heading(doc, text, level=1):
    """Add formatted heading"""
    if level == 1:
        p = doc.add_heading(text, level=1)
        p.runs[0].font.size = Pt(16)
        p.runs[0].font.bold = True
        p.runs[0].font.name = 'Times New Roman'
    elif level == 2:
        p = doc.add_heading(text, level=2)
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.bold = True
        p.runs[0].font.name = 'Times New Roman'
    return p

def add_body_text(doc, text):
    """Add body text with proper formatting"""
    p = doc.add_paragraph(text)
    p.runs[0].font.size = Pt(12)
    p.runs[0].font.name = 'Times New Roman'
    return p

# Create document
doc = Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# ===== FRONT MATTER =====
# Title Page
for _ in range(3):
    doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("AI-POWERED PERSONAL FINANCE MANAGEMENT SYSTEM")
run.font.size = Pt(16)
run.font.bold = False
run.font.name = 'Times New Roman'

doc.add_paragraph()
by_p = doc.add_paragraph("By")
by_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
by_p.runs[0].font.name = 'Times New Roman'
by_p.runs[0].font.size = Pt(12)

doc.add_paragraph()
s1 = doc.add_paragraph("Zohaib Ahmad")
s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
s1.runs[0].font.name = 'Times New Roman'
r1 = doc.add_paragraph("(Reg. No. 2022-GCUF-06910)")
r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1.runs[0].font.name = 'Times New Roman'

doc.add_paragraph()
s2 = doc.add_paragraph("Hammad Ahmad")
s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
s2.runs[0].font.name = 'Times New Roman'
r2 = doc.add_paragraph("(Reg. No. 2022-GCUF-06904)")
r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2.runs[0].font.name = 'Times New Roman'

for _ in range(5):
    doc.add_paragraph()

deg_line1 = doc.add_paragraph("BACHELOR OF SCIENCE")
deg_line1.alignment = WD_ALIGN_PARAGRAPH.CENTER
deg_line1.runs[0].font.name = 'Times New Roman'
deg_line1.runs[0].font.bold = True

deg_line2 = doc.add_paragraph("IN")
deg_line2.alignment = WD_ALIGN_PARAGRAPH.CENTER
deg_line2.runs[0].font.name = 'Times New Roman'
deg_line2.runs[0].font.bold = True

deg_line3 = doc.add_paragraph("INFORMATION TECHNOLOGY")
deg_line3.alignment = WD_ALIGN_PARAGRAPH.CENTER
deg_line3.runs[0].font.name = 'Times New Roman'
deg_line3.runs[0].font.bold = True

for _ in range(6):
    doc.add_paragraph()

dept = doc.add_paragraph("DEPARTMENT OF INFORMATION TECHNOLOGY")
dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
dept.runs[0].font.name = 'Times New Roman'
dept.runs[0].font.bold = True

uni = doc.add_paragraph("GOVERNMENT COLLEGE UNIVERSITY, FAISALABAD")
uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
uni.runs[0].font.name = 'Times New Roman'
uni.runs[0].font.bold = True

doc.add_paragraph()
month = doc.add_paragraph("June 2026")
month.alignment = WD_ALIGN_PARAGRAPH.CENTER
month.runs[0].font.name = 'Times New Roman'

doc.add_page_break()

# Declaration
decl_title = doc.add_paragraph("DECLARATION")
decl_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
decl_title.runs[0].font.name = 'Times New Roman'
decl_title.runs[0].font.size = Pt(14)
decl_title.runs[0].font.bold = True

doc.add_paragraph()
decl_text = "The work reported in this documentation was carried out by me/us under the supervision of Prof. Dr. Mariam Rehman, Department of Information Technology, Government College University, Faisalabad, Pakistan.\n\nI/We hereby declare that the title of project \"AI-Powered Personal Finance Management System\" and the contents of the documentation are the product of my own work and no part has been copied from any published source. I/we further declare that this work has not been submitted for award of any other degree/diploma."
decl_p = doc.add_paragraph(decl_text)
decl_p.runs[0].font.name = 'Times New Roman'
decl_p.runs[0].font.size = Pt(12)

doc.add_page_break()

# Certificate
cert_title = doc.add_paragraph("CERTIFICATE BY SUPERVISORY COMMITTEE")
cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
cert_title.runs[0].font.name = 'Times New Roman'
cert_title.runs[0].font.size = Pt(14)
cert_title.runs[0].font.bold = True

doc.add_paragraph()
cert_text = "We certify that the contents and form of documentation submitted by Zohaib Ahmad (Reg. 2022-GCUF-06910) and Hammad Ahmad (Reg. 2022-GCUF-06904) has been found satisfactory. We recommend it for evaluation by the External Examiner."
cert_p = doc.add_paragraph(cert_text)
cert_p.runs[0].font.name = 'Times New Roman'
cert_p.runs[0].font.size = Pt(12)

doc.add_page_break()

# Dedication
ded_title = doc.add_paragraph("DEDICATION")
ded_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
ded_title.runs[0].font.name = 'Times New Roman'
ded_title.runs[0].font.size = Pt(14)
ded_title.runs[0].font.bold = True
for _ in range(2):
    doc.add_paragraph()
ded_text = "We dedicate this Final Year Project to our families and Prof. Dr. Mariam Rehman for their invaluable support and guidance. This project represents our commitment to solving real-world problems in personal finance management."
ded_p = doc.add_paragraph(ded_text)
ded_p.runs[0].font.name = 'Times New Roman'
ded_p.runs[0].font.size = Pt(12)

doc.add_page_break()

# Acknowledgements
ack_title = doc.add_paragraph("ACKNOWLEDGEMENTS")
ack_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
ack_title.runs[0].font.name = 'Times New Roman'
ack_title.runs[0].font.size = Pt(14)
ack_title.runs[0].font.bold = True
doc.add_paragraph()
ack_text = "We express our gratitude to Prof. Dr. Mariam Rehman for her expert guidance. We thank the Department of Information Technology and our families for their constant support."
ack_p = doc.add_paragraph(ack_text)
ack_p.runs[0].font.name = 'Times New Roman'
ack_p.runs[0].font.size = Pt(12)

doc.add_page_break()

# Abstract
abs_title = doc.add_paragraph("ABSTRACT")
abs_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
abs_title.runs[0].font.name = 'Times New Roman'
abs_title.runs[0].font.size = Pt(14)
abs_title.runs[0].font.bold = True
doc.add_paragraph()
abs_text = "This project presents an AI-powered Personal Finance Management System featuring machine learning-based expense categorization, ARIMA time-series forecasting, and OCR-based receipt scanning. The system supports multi-user expense tracking, income management, goal tracking, and comprehensive reporting capabilities built using Django, scikit-learn, and modern web technologies."
abs_p = doc.add_paragraph(abs_text)
abs_p.runs[0].font.name = 'Times New Roman'
abs_p.runs[0].font.size = Pt(12)

doc.add_page_break()

# Save
doc.save("FYP_Documentation_Final_v1.docx")
print("Document created successfully!")
