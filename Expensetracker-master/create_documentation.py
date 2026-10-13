#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Complete FYP Documentation Generator
Government College University, Faisalabad
Department of Information Technology
"""

import sys
import os

print("Attempting to import python-docx...")

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    print("✓ All imports successful!")
except ImportError as e:
    print(f"✗ Import error: {e}")
    print("Installing python-docx...")
    os.system("pip install python-docx")
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_page_break(doc):
    """Add a page break"""
    doc.add_page_break()

def set_line_spacing(paragraph, spacing=1.5):
    """Set line spacing for paragraph"""
    p = paragraph._element
    pPr = p.get_or_add_pPr()
    spacing_elem = OxmlElement('w:spacing')
    spacing_elem.set(qn('w:line'), str(int(spacing * 240)))
    pPr.append(spacing_elem)

def create_title_page(doc):
    """Create Title Page"""
    for _ in range(3):
        doc.add_paragraph()
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI-POWERED PERSONAL FINANCE MANAGEMENT SYSTEM")
    run.font.size = Pt(16)
    run.font.bold = False
    run.font.name = 'Times New Roman'
    
    doc.add_paragraph()
    by = doc.add_paragraph("By")
    by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    by.runs[0].font.name = 'Times New Roman'
    by.runs[0].font.size = Pt(12)
    
    doc.add_paragraph()
    s1 = doc.add_paragraph("Zohaib Ahmad")
    s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s1.runs[0].font.name = 'Times New Roman'
    
    r1 = doc.add_paragraph("(Reg. No. 2022-GCUF-06910)")
    r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1.runs[0].font.name = 'Times New Roman'
    
    doc.add_paragraph()
    s2 = doc.add_paragraph("Hammad Ahmad")
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s2.runs[0].font.name = 'Times New Roman'
    
    r2 = doc.add_paragraph("(Reg. No. 2022-GCUF-06904)")
    r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2.runs[0].font.name = 'Times New Roman'
    
    for _ in range(5):
        doc.add_paragraph()
    
    deg1 = doc.add_paragraph("BACHELOR OF SCIENCE")
    deg1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    deg1.runs[0].font.name = 'Times New Roman'
    deg1.runs[0].font.bold = True
    
    deg2 = doc.add_paragraph("IN")
    deg2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    deg2.runs[0].font.name = 'Times New Roman'
    deg2.runs[0].font.bold = True
    
    deg3 = doc.add_paragraph("INFORMATION TECHNOLOGY")
    deg3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    deg3.runs[0].font.name = 'Times New Roman'
    deg3.runs[0].font.bold = True
    
    for _ in range(6):
        doc.add_paragraph()
    
    dept = doc.add_paragraph("DEPARTMENT OF INFORMATION TECHNOLOGY")
    dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dept.runs[0].font.name = 'Times New Roman'
    dept.runs[0].font.bold = True
    
    uni = doc.add_paragraph("GOVERNMENT COLLEGE UNIVERSITY, FAISALABAD")
    uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    uni.runs[0].font.name = 'Times New Roman'
    uni.runs[0].font.bold = True
    
    doc.add_paragraph()
    month = doc.add_paragraph("June 2026")
    month.alignment = WD_ALIGN_PARAGRAPH.CENTER
    month.runs[0].font.name = 'Times New Roman'
    
    add_page_break(doc)

def create_declaration(doc):
    """Create Declaration Page"""
    title = doc.add_paragraph("DECLARATION")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.name = 'Times New Roman'
    title.runs[0].font.size = Pt(14)
    title.runs[0].font.bold = True
    
    doc.add_paragraph()
    
    text = (
        "The work reported in this documentation was carried out by me/us under the supervision of "
        "Prof. Dr. Mariam Rehman, Department of Information Technology, Government College University, "
        "Faisalabad, Pakistan.\n\n"
        "I/We hereby declare that the title of project \"AI-Powered Personal Finance Management System\" "
        "and the contents of the documentation are the product of my own work and no part has been copied from "
        "any published source (except the references, standard mathematical or genetic models/equations/formulas/"
        "protocols etc.). I/we further declare that this work has not been submitted for award of any other degree/"
        "diploma. The University may take action if the information provided is found inaccurate at any stage."
    )
    
    p = doc.add_paragraph(text)
    set_line_spacing(p, 1.5)
    p.runs[0].font.name = 'Times New Roman'
    p.runs[0].font.size = Pt(12)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Signature of Student"
    table.rows[1].cells[0].text = "Zohaib Ahmad (2022-GCUF-06910)"
    table.rows[0].cells[1].text = "Signature of Student"
    table.rows[1].cells[1].text = "Hammad Ahmad (2022-GCUF-06904)"
    
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.runs[0].font.name = 'Times New Roman'
                paragraph.runs[0].font.size = Pt(11)
    
    add_page_break(doc)

def create_abstract(doc):
    """Create Abstract"""
    title = doc.add_paragraph("ABSTRACT")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.name = 'Times New Roman'
    title.runs[0].font.size = Pt(14)
    title.runs[0].font.bold = True
    
    doc.add_paragraph()
    
    text = (
        "Personal finance management is a critical aspect of financial well-being, yet many individuals struggle "
        "with effective expense tracking, budgeting, and financial planning. This Final Year Project presents an "
        "AI-powered Personal Finance Management System designed to address these challenges.\n\n"
        "The system leverages cutting-edge technologies including machine learning for intelligent category prediction, "
        "time-series forecasting using ARIMA models, and Optical Character Recognition (OCR) for automated receipt scanning.\n\n"
        "Key Features:\n"
        "• Multi-user expense and income tracking with PKR currency support\n"
        "• Intelligent category prediction using TF-IDF + Random Forest (89% accuracy)\n"
        "• 30-day expense forecasting using ARIMA(5,1,0) model\n"
        "• Automated receipt scanning and field extraction using EasyOCR\n"
        "• Savings goal management with progress tracking\n"
        "• Comprehensive reporting (PDF, CSV, Excel export)\n"
        "• Secure authentication and multi-user data isolation\n"
        "• Interactive dashboard with Chart.js visualizations\n\n"
        "Built on Django 5.1.1 with SQLite (development) and PostgreSQL (production), the system implements industry-standard "
        "security practices. Performance metrics show sub-second dashboard response times and efficient multi-user handling.\n\n"
        "This documentation provides complete technical specifications, system design, implementation details, testing strategies, "
        "and user interface documentation for deployment and future enhancement."
    )
    
    p = doc.add_paragraph(text)
    set_line_spacing(p, 1.5)
    p.runs[0].font.name = 'Times New Roman'
    p.runs[0].font.size = Pt(12)
    
    add_page_break(doc)

def main():
    """Main function"""
    print("Starting FYP Documentation Generation...")
    
    # Create document
    doc = Document()
    
    # Set default margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    print("Creating Title Page...")
    create_title_page(doc)
    
    print("Creating Declaration...")
    create_declaration(doc)
    
    print("Creating Abstract...")
    create_abstract(doc)
    
    # Save document
    filename = "FYP_Documentation_Expense_Tracker.docx"
    doc.save(filename)
    
    print(f"\n✓ Document created successfully!")
    print(f"✓ Filename: {filename}")
    print(f"✓ Location: {os.path.abspath(filename)}")
    print(f"✓ File size: {os.path.getsize(filename)} bytes")

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"✗ Error: {e}")
        import traceback
        traceback.print_exc()
